# -*- coding: utf-8 -*-
"""
Created on Thu Jul 30 12:12:45 2026

@author: hp
"""

import cv2
import numpy as np
import os
from datetime import datetime
from keras.models import load_model
from openpyxl import Workbook, load_workbook

# ----------------------------------------------------
# 1) LOAD FACE DETECTOR AND TRAINED MODEL
# ----------------------------------------------------
classifier = cv2.CascadeClassifier(
    cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
)

model = load_model("Final_model.h5")

# Open Webcam
cap = cv2.VideoCapture(0)

# ----------------------------------------------------
# 2) LABELS
# ----------------------------------------------------
labels = ["Mayank"]      # Add more names if required


def get_pred_label(pred):
    return labels[pred]


# ----------------------------------------------------
# 3) PREPROCESS
# ----------------------------------------------------
def preprocess(img):
    img = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    img = cv2.resize(img, (100, 100))
    img = cv2.equalizeHist(img)
    img = img.reshape(1, 100, 100, 1)
    img = img / 255.0
    return img


# ----------------------------------------------------
# 4) ATTENDANCE VARIABLES
# ----------------------------------------------------
attendance_started = False
attendance_file = ""
recorded_names = set()


# ----------------------------------------------------
# 5) START ATTENDANCE
# ----------------------------------------------------
def start_attendance():
    global attendance_started, attendance_file, recorded_names

    recorded_names.clear()

    filename = datetime.now().strftime("Attendance_%Y-%m-%d.xlsx")
    attendance_file = filename

    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        ws.append([
            "Name",
            "Date",
            "Start Time",
            "End Time"
        ])

        wb.save(filename)

    attendance_started = True
    print("\nAttendance Started")


# ----------------------------------------------------
# 6) SAVE START TIME
# ----------------------------------------------------
def save_attendance(name):

    wb = load_workbook(attendance_file)
    ws = wb.active

    # Already recorded
    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            wb.close()
            return

    now = datetime.now()

    ws.append([
        name,
        now.strftime("%Y-%m-%d"),
        now.strftime("%H:%M:%S"),
        ""
    ])

    wb.save(attendance_file)
    wb.close()

    print(f"{name} Marked Present")


# ----------------------------------------------------
# 7) END ATTENDANCE
# ----------------------------------------------------
def end_attendance():

    global attendance_started

    if not attendance_started:
        return

    wb = load_workbook(attendance_file)
    ws = wb.active

    end_time = datetime.now().strftime("%H:%M:%S")

    for row in ws.iter_rows(min_row=2):

        if row[3].value is None or row[3].value == "":
            row[3].value = end_time

    wb.save(attendance_file)
    wb.close()

    attendance_started = False

    print("\nAttendance Finished")
    print("Excel Saved:", attendance_file)


# ----------------------------------------------------
# 8) MAIN LOOP
# ----------------------------------------------------
while True:

    ret, frame = cap.read()

    if not ret:
        break

    faces = classifier.detectMultiScale(
        frame,scaleFactor=1.3,minNeighbors=5
    )

    for (x, y, w, h) in faces:

        face = frame[y:y+h, x:x+w]

        cv2.rectangle(
            frame,(x, y),(x+w, y+h),(255, 0, 0),2
        )

        prediction = np.argmax(
            model.predict(
                preprocess(face),verbose=0
            )
        )

        name = get_pred_label(prediction)

        cv2.putText(
            frame,name,(x, y-10),
            cv2.FONT_HERSHEY_SIMPLEX,0.8,(0, 255, 0),2
        )

        # Save attendance only once
        if attendance_started:
            if name not in recorded_names:
                save_attendance(name)
                recorded_names.add(name)

    # -------------------------
    # Draw Buttons
    # -------------------------
    cv2.rectangle(frame, (10, 10), (180, 55), (0, 180, 0), -1)
    cv2.putText(
        frame,"S - START",(25, 40),
        cv2.FONT_HERSHEY_SIMPLEX,0.8,(255, 255, 255),2
    )

    cv2.rectangle(frame, (200, 10), (360, 55), (0, 0, 255), -1)
    cv2.putText(
        frame,"E - END",
        (225, 40),cv2.FONT_HERSHEY_SIMPLEX,0.8,(255, 255, 255),2
    )

    cv2.imshow("Face Attendance System", frame)

    key = cv2.waitKey(1) & 0xFF

    # Start Attendance
    if key == ord('s'):

        if not attendance_started:
            start_attendance()

    # End Attendance
    elif key == ord('e'):

        end_attendance()

        print("\nClosing Application...")

        break

    # Quit
    elif key == ord('q'):
        break


# ----------------------------------------------------
# 9) RELEASE
# ----------------------------------------------------
cap.release()
cv2.destroyAllWindows()